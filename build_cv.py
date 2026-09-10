# -*- coding: utf-8 -*-
"""Generate a clean, interview-friendly CV in the format of GeneralCV_RSI_processed_rsiindia.docx."""
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_TAB_ALIGNMENT, WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

TNR = "Times New Roman"
BLUE = "2E74B5"
NAVY = "1F3864"
SILVER = "C0C0C0"

doc = Document()

sec = doc.sections[0]
sec.page_width = Inches(8.5)
sec.page_height = Inches(11)
sec.top_margin = Inches(0.5)
sec.bottom_margin = Inches(0.5)
sec.left_margin = Inches(0.5)
sec.right_margin = Inches(0.5)

normal = doc.styles["Normal"]
normal.font.name = TNR
normal.font.size = Pt(10)


def set_shading(paragraph, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    paragraph._p.get_or_add_pPr().append(shd)


def set_bottom_border(paragraph, sz, color, space="1"):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(sz))
    bottom.set(qn("w:space"), space)
    bottom.set(qn("w:color"), color)
    pBdr.append(bottom)
    pPr.append(pBdr)


def sp_after(paragraph, pts):
    paragraph.paragraph_format.space_after = Pt(pts)
    paragraph.paragraph_format.space_before = Pt(0)


def add_run(p, text, bold=False, size=10, color=None, name=TNR):
    r = p.add_run(text)
    r.bold = bold
    r.font.name = name
    r.font.size = Pt(size)
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    return r


def name_banner(text):
    p = doc.add_paragraph()
    sp_after(p, 0)
    set_shading(p, NAVY)
    add_run(p, text, bold=True, size=16, color="FFFFFF")


def title_line(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    set_bottom_border(p, 12, SILVER, space="0")
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.space_before = Pt(2)
    add_run(p, text, bold=True, size=12)


def section(text):
    p = doc.add_paragraph()
    sp_after(p, 2)
    set_bottom_border(p, 18, BLUE, space="1")
    add_run(p, "|| " + text, bold=True, size=11, color=BLUE)


def bullet(text=None, runs=None):
    p = doc.add_paragraph(style="List Bullet")
    sp_after(p, 0)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if runs:
        for t, b in runs:
            add_run(p, t, bold=b, size=10)
    else:
        add_run(p, text, size=10)


def plain(runs, space=0, justify=True):
    p = doc.add_paragraph()
    sp_after(p, space)
    if justify:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for t, b in runs:
        add_run(p, t, bold=b, size=10)


def blank():
    p = doc.add_paragraph()
    sp_after(p, 0)


def project_header(title, role, duration):
    pt = doc.add_paragraph()
    sp_after(pt, 0)
    add_run(pt, title, bold=True, size=10)
    pr = doc.add_paragraph()
    sp_after(pr, 0)
    pr.paragraph_format.tab_stops.add_tab_stop(Inches(7.5), WD_TAB_ALIGNMENT.RIGHT)
    add_run(pr, "Role:", bold=True, size=10)
    add_run(pr, " " + role, size=10)
    add_run(pr, "\t", size=10)
    add_run(pr, "Duration:", bold=True, size=10)
    add_run(pr, " " + duration, size=10)


def description(text):
    plain([("Description:", True), (" " + text, False)], space=0)


def responsibilities():
    plain([("Responsibilities: ", True)], space=0)


def tech_stack(text):
    bullet(runs=[("Tech Stack", True), (": " + text, False)])


# ============================ CONTENT ============================
name_banner("Bali Mohan Rao")
title_line("Associate Data Engineer | Cloud Data Engineering & BI Modernization")

# ---- Profile Summary ----
section("Profile Summary")
bullet("Data Engineer at R Systems International with close to a year of hands-on experience building data migration tools, cloud ETL pipelines, and BI modernization platforms.")
bullet("Databricks Certified Data Engineer Associate and Databricks Certified Generative AI Engineer Associate (2026), with strong fundamentals in Apache Spark, Delta Lake, and data warehousing.")
bullet("Delivered 6 end-to-end projects across client and internal work, including an Azure Analysis Services to Power BI migration tool that reduced manual effort by around 60% (earned an R Systems SPOT Award) and a Domo to Snowflake platform for a Fortune 500 client.")
bullet("Strong in Python, SQL, Snowflake, Databricks, Apache Spark, Azure, Microsoft Fabric, and Power BI, with a focus on automating data pipelines and reducing manual work.")
bullet("1st Place winner at BeEXIQO Code-AI-Thon 2026. Build projects using an AI-first approach with Cursor and modern LLMs, in line with my organization's AI-first engineering practices.")
blank()

# ---- Education & Certifications ----
section("Education & Certifications")
bullet(runs=[("Databricks Certified Data Engineer Associate", True), (" | May 2026 (Valid to May 2028)", False)])
bullet(runs=[("Databricks Certified Generative AI Engineer Associate", True), (" | June 2026 (Valid to June 2028)", False)])
bullet("Bachelor of Technology - Computer Science and Engineering")
plain([("Lovely Professional University, Punjab, India | CGPA: 8.05 | Aug 2022 - Jun 2026", False)], space=0)
bullet("Cloud Computing Certification - NPTEL | Jan 2025 - May 2025")
bullet("Supervised Machine Learning - Coursera | Sep 2024 - Oct 2024")
blank()

# ---- Technical Skills ----
section("Technical Skills")
skills = [
    ("Cloud & Data Platforms", "Databricks (Certified), Snowflake, Azure Analysis Services, Microsoft Fabric, Azure Data Factory, Power BI Service, Domo"),
    ("Languages & Query", "Python, SQL (T-SQL, Snowflake SQL, Spark SQL), DAX, TypeScript, JavaScript"),
    ("Data Engineering & ETL", "Apache Spark, Delta Lake, Medallion Architecture, ETL Pipelines, Metadata Extraction, Schema Automation, Data Migration, REST API Integration, dbt"),
    ("Frameworks & Tools", "Flask, Node.js, pandas, SQLAlchemy, Redis, Docker, Git, Cursor (AI-assisted development)"),
    ("AI / Generative AI", "OpenAI, Anthropic, Azure OpenAI, Groq, RAG, Prompt Engineering, Multi-Agent Systems, Neo4j"),
    ("BI & Analytics", "Power BI (Semantic Models, DAX, TMDL), SSMS, Report Rebinding, KPI Dashboards, Domo"),
]
table = doc.add_table(rows=len(skills), cols=2)
table.autofit = False
for i, (cat, val) in enumerate(skills):
    c0 = table.rows[i].cells[0]
    c1 = table.rows[i].cells[1]
    c0.width = Inches(2.1)
    c1.width = Inches(5.4)
    p0 = c0.paragraphs[0]
    sp_after(p0, 0)
    add_run(p0, cat, bold=True, size=10)
    p1 = c1.paragraphs[0]
    sp_after(p1, 0)
    add_run(p1, val, size=10)
blank()

# ---- Work Experience ----
section("Work Experience")
plain([("R Systems International", True), ("   |   Associate Data Engineer   |   Jun 2025 - Present   |   Pune, India", False)], space=2)

# Project 1 - MetricAI
project_header("Project#1: MetricAI - AAS to Power BI Semantic Migration Accelerator (Phase 1 & 2)",
               "Associate Data Engineer", "2025")
description("An AI-powered tool to migrate Azure Analysis Services (AAS) models to Databricks-backed Power BI, combining automation scripts with an AI-driven Flask application.")
responsibilities()
bullet("Built a set of automation tools for complexity assessment, metadata extraction, and report rebinding that cut manual migration effort by around 60%.")
bullet("Developed a Flask application that reads AAS metadata, maps DAX measure dependencies, and automatically generates Databricks SQL views.")
bullet("Added AI-assisted logic to standardize KPIs and generate Power BI semantic models automatically.")
bullet("Created a validation step to compare source and generated outputs, with a review-and-approval workflow and export to Excel, SQL, and JSON.")
bullet("Converted Power BI datasets from Import mode to DirectQuery connected to Databricks and optimized DAX measures for better performance.")
tech_stack("Python, Flask, Azure Analysis Services, Databricks SQL, Power BI, TMDL, DAX, SQL, OpenAI API")
blank()

# Project 2 - Domo to Snowflake
project_header("Project#2: Domo to Snowflake & Power BI Migration Accelerator",
               "Associate Data Engineer", "2026")
description("An internal tool to automate migration of Domo dashboards, datasets, and ETL pipelines into a Snowflake and Power BI setup for a Fortune 500 client.")
responsibilities()
bullet("Used Domo APIs to automatically list and extract 100+ dashboards and ETL pipelines, reducing migration time from weeks to hours.")
bullet("Built a module that automatically creates Snowflake tables based on the extracted Domo metadata.")
bullet("Automated Power BI semantic model creation and dashboard publishing using the Power BI REST API.")
bullet("Developed a batch data migration pipeline (SQL Server to Snowflake) with row-count checks to ensure complete and accurate data transfer.")
bullet("Mapped Domo chart types to matching Power BI visuals for automated report rebuilding.")
tech_stack("Python, Flask, Domo APIs, Snowflake, Power BI REST API, SQL Server, pandas")
blank()

# Project 3 - MetaSphere
project_header("Project#3: MetaSphere - Enterprise AI Data Intelligence Platform",
               "Full-Stack / AI Data Engineer", "2026")
description("A full-stack AI data platform that connects to existing databases, understands their structure, and uses AI agents to answer business questions in plain English - running queries directly at the source without moving data.")
responsibilities()
bullet("Built a full-stack platform (React and Node.js) that connects to SQL Server, PostgreSQL, Snowflake, and Databricks and queries them directly, without copying data.")
bullet("Designed a central AI system that understands a user's question, plans the steps, and routes them to the right tools in a controlled and fully logged way.")
bullet("Built a feature that splits one question into queries across multiple databases, runs them in parallel, and combines the results using fuzzy matching to link related records.")
bullet("Added role-based access control with four data-sensitivity levels, a complete audit trail, and support for multiple AI providers with secure key storage.")
bullet("Implemented a Neo4j knowledge graph for data lineage, document search using RAG, and a multi-stage review process to check AI answer quality.")
tech_stack("React, TypeScript, Node.js, Express, PostgreSQL, Redis, Neo4j, SQL Server, Snowflake, OpenAI/Anthropic, Docker")
blank()

# Project 4 - Anything to Snowflake
project_header("Project#4: Anything to Snowflake - Enterprise AI Migration Accelerator",
               "Associate Data Engineer", "2026")
description("A multi-agent platform to migrate different data sources into Snowflake using a layered Bronze/Silver/Gold architecture with built-in governance and monitoring.")
responsibilities()
bullet("Designed a modular system of 12 agents (for source discovery, table creation, governance, drift detection, and self-healing) that communicate through an event-driven design using Redis.")
bullet("Implemented a Bronze/Silver/Gold layered architecture on Snowflake with automatic table creation and incremental data loading using Snowflake Streams and Tasks.")
bullet("Built a governance layer that automatically detects sensitive data (PII/PCI), applies masking, controls row-level access, and tracks data lineage end to end.")
bullet("Added a self-healing feature that automatically detects and retries failed steps based on defined rules.")
bullet("Deployed the platform with Docker and a secure REST API (40+ endpoints) with role-based access.")
tech_stack("Python, Flask, PostgreSQL, Redis, Snowflake, Docker, SQLAlchemy, JWT/RBAC")
blank()

# Project 5 - DataBridge Pro
project_header("Project#5: DataBridge Pro - SQL Server & Excel to Cloud Migration Platform",
               "Associate Data Engineer", "2026")
description("A tool to check migration readiness and run live migrations from SQL Server and Excel into Snowflake, Microsoft Fabric, and Power BI.")
responsibilities()
bullet("Built a single tool supporting SQL Server to Snowflake, SQL Server to Microsoft Fabric, Excel to Power BI, and SQL Server stored procedures to dbt models.")
bullet("Created an assessment feature that analyzes SQL Server stored procedures, views, and table relationships and scores how ready they are to move to Snowflake or Fabric.")
bullet("Implemented live data migration in batches with secure authentication for both Snowflake and Microsoft Fabric.")
bullet("Built a 9-step Excel analysis that detects KPIs, charts, and formulas and gives a Power BI readiness score.")
bullet("Developed an Excel-to-DAX formula converter and an automated Power BI project (PBIP) generator.")
tech_stack("Python, Flask, pandas, openpyxl, Snowflake, Microsoft Fabric, Azure Data Factory, dbt")
blank()

# Project 6 - TelCortex AI
project_header("Project#6: TelCortex AI - Real-Time Telecom Intelligence Platform",
               "AI Engineer (Hackathon)", "BeEXIQO Code-AI-Thon 2026")
description("An AI platform with 5 agents that analyzes telecom call records in real time to detect fraud, network issues, revenue leakage, service-quality problems, and customer churn - winner of 1st Place at BeEXIQO Code-AI-Thon 2026.")
responsibilities()
bullet("Won 1st Place at BeEXIQO Code-AI-Thon 2026 by building a 5-agent platform (fraud, network, revenue, service quality, and churn) that processes call records in real time.")
bullet("Built a central engine that combines signals from all agents to decide the right action - from raising an alert to creating a ticket, dispatching a field engineer, or sending a retention offer.")
bullet("Added automated workflows with human approval for important actions, plus clear explanations for each alert (confidence level and likely root cause).")
bullet("Built real-time data streaming using WebSockets and a chat interface to ask questions about telecom data in plain English.")
bullet("Packaged the full platform using Docker across 5 services.")
tech_stack("Python, Flask, Celery, Redis, PostgreSQL, Next.js, Groq (Llama 3.1), Docker")
blank()

# ---- Awards & Achievements ----
section("Awards & Achievements")
bullet(runs=[("R Systems SPOT Award - Outstanding Performance (2025):", True), (" Recognized for outstanding contribution to the AAS to Power BI migration tool, which reduced manual effort by around 60%.", False)])
bullet(runs=[("1st Place - BeEXIQO Code-AI-Thon 2026 (Team Predators):", True), (" Won the hackathon by building TelCortex AI, a real-time telecom intelligence platform with 5 AI agents.", False)])

out = r"C:\Users\bali\Desktop\Mohan\CVs\BaliMohanRao_CV_2026.docx"
doc.save(out)
print("SAVED:", out)
print("paragraphs:", len(doc.paragraphs), "tables:", len(doc.tables))
